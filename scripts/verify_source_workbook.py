"""Read-only verification and extraction for the Hongxing source workbook.

The workbook is a business source of truth, so this module deliberately does
not write to it or silently repair it. It keeps source values and diagnostics
alongside the parsed result. ``extract`` retains the ``cases`` shape used by
the isolated acceptance runner while adding v2 metadata for the flow, HR
summary sheets, shared formulas, and source anomalies.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries


NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
SCORER_ALIASES = {
    "PM": "PM", "项目经理": "PM", "项目经理（PM）": "PM",
    "CTO": "CTO", "部门负责人": "DEPARTMENT_MANAGER",
    "部门负责人（E层）": "DEPARTMENT_MANAGER", "部门负责人(E层)": "DEPARTMENT_MANAGER",
    "L层Leader": "DEPARTMENT_MANAGER", "L 层 Leader": "DEPARTMENT_MANAGER",
    "产品经理": "PRODUCT_MANAGER", "产品经理（PD）": "PRODUCT_MANAGER", "产品经理(PD)": "PRODUCT_MANAGER",
}
DEPARTMENT_FALLBACKS = {
    "01-ID工业设计": "ID 工业设计部", "02-结构部门": "结构部门", "03-电子部门": "电子部门",
    "04-软件应用开发": "软件应用开发部", "05-算法开发": "算法开发部", "06-运动控制": "运控部",
    "07-测试部门": "测试部门", "08-嵌入式开发": "嵌入式开发部", "09-产品部门": "产品部门",
}


def rounded(value: Any, places: int) -> Decimal:
    return Decimal(str(value)).quantize(Decimal(10) ** -places, rounding=ROUND_HALF_UP)


def grade(score: Any) -> str:
    score = Decimal(str(score))
    return next((label for limit, label in [(90, "S"), (80, "A"), (70, "B"), (60, "C")] if score >= limit), "D")


def _formula_text(value: Any) -> str | None:
    """Return a JSON-safe formula with a leading ``=``."""
    if hasattr(value, "text"):
        value = value.text
    if not isinstance(value, str) or not value:
        return None
    return value if value.startswith("=") else "=" + value


def _cell_value(cell: Any) -> Any:
    value = cell.value
    if hasattr(value, "text"):
        return {"formula": _formula_text(value), "ref": getattr(value, "ref", None)}
    return value


def _translate_formula(formula: str, source: str, target: str) -> str:
    try:
        return Translator(formula, origin=source).translate_formula(target)
    except Exception:
        source_row = int(re.search(r"\d+", source).group())
        target_row = int(re.search(r"\d+", target).group())
        return re.sub(r"(?<![A-Z$])([A-Z]{1,3})(\d+)", lambda m: m.group(1) + str(int(m.group(2)) + target_row - source_row), formula)


def _worksheet_paths(path: str | Path) -> dict[str, str]:
    """Map workbook sheet names to worksheet XML paths."""
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rels = {r.get("Id"): r.get("Target") for r in rels_root.findall(f"{{{PKG_REL_NS}}}Relationship")}
        output: dict[str, str] = {}
        for sheet in workbook.findall("s:sheets/s:sheet", NS):
            target = rels.get(sheet.get(f"{{{REL_NS}}}id"), "").lstrip("/")
            output[sheet.get("name", "")] = posixpath.normpath(posixpath.join("xl", target))
        return output


def _formula_topology(path: str | Path) -> dict[str, dict[str, Any]]:
    """Extract formulas and expand shared-formula followers from raw XML."""
    topology: dict[str, dict[str, Any]] = {}
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        for sheet_name, sheet_path in _worksheet_paths(path).items():
            if sheet_path not in names:
                continue
            root = ET.fromstring(archive.read(sheet_path))
            cells: dict[str, dict[str, Any]] = {}
            shared: dict[str, dict[str, Any]] = {}
            for cell in root.findall("s:sheetData/s:row/s:c", NS):
                coord = cell.get("r")
                formula_node = cell.find("s:f", NS)
                if coord is None:
                    continue
                attrs = dict(formula_node.attrib) if formula_node is not None else {}
                raw = formula_node.text if formula_node is not None else None
                cells[coord] = {"formula": ("=" + raw) if raw else None, "source_formula": ("=" + raw) if raw else None,
                                "formula_attributes": attrs, "shared_expected": False}
                if attrs.get("t") == "shared" and raw:
                    shared[attrs.get("si")] = {"origin": coord, "formula": "=" + raw, "ref": attrs.get("ref")}
            gaps: list[dict[str, str]] = []
            groups: list[dict[str, Any]] = []
            for si, master in shared.items():
                ref = master.get("ref")
                if not ref:
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(ref)
                expected: list[str] = []
                for row in range(min_row, max_row + 1):
                    for col in range(min_col, max_col + 1):
                        target = f"{get_column_letter(col)}{row}"
                        expected.append(target)
                        if target not in cells:
                            gaps.append({"cell": target, "shared_index": str(si), "range": ref, "reason": "missing_cell"})
                            continue
                        cells[target]["shared_expected"] = True
                        cells[target]["expanded_formula"] = _translate_formula(master["formula"], master["origin"], target)
                        if not cells[target].get("formula"):
                            cells[target]["formula"] = cells[target]["expanded_formula"]
                groups.append({"shared_index": str(si), "origin": master["origin"], "range": ref, "expanded_cells": expected})
            topology[sheet_name] = {"cells": cells, "shared_formula_groups": groups, "gaps": gaps}
    return topology


def _find_department(sheet: Any) -> dict[str, Any]:
    for row in sheet.iter_rows():
        for cell in row:
            if str(cell.value).strip() != "部门":
                continue
            for col in range(cell.column + 1, sheet.max_column + 1):
                value = sheet.cell(cell.row, col).value
                if value is not None and str(value).strip():
                    return {"value": str(value).strip(), "source_cell": sheet.cell(cell.row, col).coordinate, "label_cell": cell.coordinate, "inferred": False}
            return {"value": None, "source_cell": None, "label_cell": cell.coordinate, "inferred": True}
    return {"value": None, "source_cell": None, "label_cell": None, "inferred": True}


def _parse_flow(sheet: Any) -> dict[str, Any]:
    steps = []
    if sheet is None:
        return {"sheet": None, "scope": "", "steps": []}
    for row in range(1, sheet.max_row + 1):
        step = sheet.cell(row, 1).value
        if not isinstance(step, (int, float)):
            continue
        steps.append({"step": int(step), "name": sheet.cell(row, 2).value, "responsibility": sheet.cell(row, 4).value,
                      "requirements": sheet.cell(row, 5).value, "deadline": sheet.cell(row, 9).value,
                      "output": sheet.cell(row, 10).value, "source_row": row})
    return {"sheet": sheet.title, "scope": sheet.cell(2, 1).value or "", "steps": steps}


def _extract_source_data(sheet: Any, template_index: int) -> dict[str, Any]:
    title = sheet.title
    position = "部门负责人" if title.endswith("部门负责人") else "工程师" if title.endswith("工程师") else ""
    level = "L" if position == "部门负责人" else "E" if position == "工程师" else None
    base = title[: -(len(position) + 1)] if position else title
    department_meta = _find_department(sheet)
    raw_department = department_meta.get("value")
    department = raw_department or DEPARTMENT_FALLBACKS.get(base) or base
    department_meta.update({"raw_value": raw_department, "value": department, "fallback": raw_department is None})
    headers = [cell.row for row in sheet.iter_rows() for cell in row if cell.value == "考核指标"]
    issues: list[dict[str, Any]] = []
    if len(headers) != 1:
        issues.append({"code": "template_header_count", "headers": headers})
        if not headers:
            return {"sheet": title, "department": department, "position": position, "level": level, "dimensions": [], "items": [],
                    "source_issues": issues, "formula_issues": [], "metadata_issues": ["missing_indicator_header"],
                    "source_metadata": {"department": department_meta, "position_raw": position, "level": level}, "composite_sum": 0.0,
                    "expected_total": 0.0, "expected_grade": grade(0)}
    header = headers[0]
    dimensions: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    formula_issues: list[dict[str, Any]] = []
    metadata_issues: list[dict[str, Any]] = []
    for row in range(header + 1, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if str(label).startswith(("综合权重合计", "总分", "绩效等级")):
            break
        name = sheet.cell(row, 4).value
        item_weight = sheet.cell(row, 5).value
        if not name or not isinstance(item_weight, (int, float)):
            continue
        dimension_name = sheet.cell(row, 2).value
        if dimension_name is not None:
            dimensions.append({"name": str(dimension_name), "weight": sheet.cell(row, 3).value,
                               "source_cell": f"C{row}", "sort_order": len(dimensions), "items": []})
        elif not dimensions:
            metadata_issues.append({"code": "orphan_indicator", "source_row": row, "name": name})
            continue
        current = dimensions[-1]
        formula = _formula_text(sheet.cell(row, 6).value)
        dim_row = int(re.search(r"\d+", current["source_cell"]).group())
        canonical_formula = f"=ROUND($C${dim_row}*E{row}/100,1)"
        if not formula or formula.replace(" ", "") != canonical_formula.replace(" ", ""):
            formula_issues.append({"code": "composite_formula_mismatch", "source_cell": f"F{row}", "actual": formula, "expected": canonical_formula})
        scorer_label = str(sheet.cell(row, 7).value or "").strip()
        scorer_type = SCORER_ALIASES.get(scorer_label, "UNKNOWN")
        if scorer_type == "UNKNOWN":
            metadata_issues.append({"code": "unknown_scorer", "source_cell": f"G{row}", "value": scorer_label})
        dimension_value = current.get("weight")
        if not isinstance(dimension_value, (int, float)):
            metadata_issues.append({"code": "missing_dimension_weight", "source_cell": current.get("source_cell"), "value": dimension_value})
            dimension_value = 0
        composite = rounded(Decimal(str(dimension_value)) * Decimal(str(item_weight)) / 100, 1)
        original_score = sheet.cell(row, 8).value
        # H is a source score/result column, not a template fixture.  A blank
        # cell must remain blank so diagnostics never manufacture a score for
        # an indicator that has not actually been assessed.
        score = original_score if isinstance(original_score, (int, float)) else None
        guidance = sheet.cell(row, 10).value or ""
        data_source_match = re.search(r"数据来源[：:]?\s*(.*)$", str(guidance), re.S)
        data_source = data_source_match.group(1).strip() if data_source_match else ""
        contribution_formula = _formula_text(sheet.cell(row, 9).value)
        contribution = rounded(composite * Decimal(str(score)) / 100, 2) if score is not None else None
        item = {"name": str(name), "item_weight": item_weight, "scorer_types": [scorer_type], "required": True,
                "scoring_guidance": str(guidance), "data_source": data_source, "order": len(items), "source_row": row,
                "source_scorer_label": scorer_label, "composite_formula": formula, "contribution_formula": contribution_formula,
                "composite_weight": float(composite), "score": score,
                "score_origin": "source" if isinstance(original_score, (int, float)) else "blank",
                "source_score": original_score, "source_score_cell": f"H{row}",
                "contribution": float(contribution) if contribution is not None else None}
        current["items"].append(item)
        items.append(item)
    for dimension in dimensions:
        total = sum(Decimal(str(item["item_weight"])) for item in dimension["items"])
        if total != 100:
            issues.append({"code": "dimension_item_weight_sum", "dimension": dimension["name"], "weight_sum": float(total),
                           "cells": [f"E{i['source_row']}" for i in dimension["items"]]})
    composite_sum = sum(Decimal(str(item["composite_weight"])) for item in items)
    scored_items = [item for item in items if item["contribution"] is not None]
    total = sum(Decimal(str(item["contribution"])) for item in scored_items)
    score_complete = len(scored_items) == len(items)
    return {"sheet": title, "department": department, "position": position, "level": level, "dimensions": dimensions,
            "items": items, "source_issues": issues, "formula_issues": formula_issues, "metadata_issues": metadata_issues,
            "source_metadata": {"department": department_meta, "position_raw": position, "level": level, "indicator_header_row": header},
            "composite_sum": float(composite_sum),
            "expected_total": float(total) if score_complete else None,
            "expected_grade": grade(total) if score_complete else None,
            "score_complete": score_complete}


def _parse_hr(sheet: Any, topology: dict[str, Any]) -> dict[str, Any]:
    quarterly = "季度" in sheet.title
    detail_start, detail_end = (5, 34) if quarterly else (5, 33)
    formula_gaps = []
    for row in range(detail_start, detail_end + 1):
        # A is only the pre-filled sequence number.  The source workbook has
        # numbered rows ready for future entries, so it cannot establish that
        # a real assessment exists.  B:H hold the business inputs; preserve 0
        # as a populated input and treat whitespace-only text as blank.
        input_cells = [
            sheet.cell(row, col).coordinate for col in range(2, 9)
            if sheet.cell(row, col).value is not None
            and (not isinstance(sheet.cell(row, col).value, str) or sheet.cell(row, col).value.strip())
        ]
        is_real_data_row = bool(input_cells)
        if sheet.cell(row, 1).value is None and not is_real_data_row:
            continue
        for col in (9, 10):
            coord = f"{get_column_letter(col)}{row}"
            if not _formula_text(sheet.cell(row, col).value):
                formula_gaps.append({"cell": coord, "row": row, "column": get_column_letter(col),
                                     "reason": "detail_row_formula_missing", "is_real_data_row": is_real_data_row,
                                     "severity": "blocking" if is_real_data_row else "warning", "input_cells": input_cells})
    expected_period_columns = list(range(4, 8)) if quarterly else list(range(4, 16))
    annual_weight_row = 39 if quarterly else 38
    annual_formula_issues = []
    annual_formula_column = "K" if quarterly else "P"
    for row in range(40, 50):
        formula = _formula_text(sheet[f"{annual_formula_column}{row}"].value)
        if not formula:
            continue
        expected_last = get_column_letter(expected_period_columns[-1])
        expected_weight = f"${get_column_letter(expected_period_columns[0])}${annual_weight_row}:${expected_last}${annual_weight_row}"
        expected_values = f"{get_column_letter(expected_period_columns[0])}{row}:{expected_last}{row}"
        if expected_weight not in formula or expected_values not in formula:
            annual_formula_issues.append({"cell": f"{annual_formula_column}{row}", "formula": formula,
                                           "expected_period_columns": [get_column_letter(c) for c in expected_period_columns],
                                           "reason": "annual_formula_scope"})
    cells = {}
    for row in sheet.iter_rows():
        for cell in row:
            value = _cell_value(cell)
            formula = _formula_text(cell.value)
            if value is not None or formula:
                xml_cell = topology.get("cells", {}).get(cell.coordinate, {})
                cells[cell.coordinate] = {
                    "value": value,
                    "formula": formula,
                    "source_formula": xml_cell.get("source_formula"),
                    "expanded_formula": xml_cell.get("expanded_formula") or formula,
                    "formula_attributes": xml_cell.get("formula_attributes", {}),
                    "formula_status": "expanded_shared" if xml_cell.get("shared_expected") else ("formula" if formula else "value"),
                }
    return {"sheet": sheet.title, "kind": "quarterly" if quarterly else "monthly", "detail_range": f"A{detail_start}:J{detail_end}",
            "annual_formula_column": annual_formula_column, "annual_weight_row": annual_weight_row,
            "period_columns": [get_column_letter(c) for c in expected_period_columns], "formula_gaps": formula_gaps,
            "annual_formula_issues": annual_formula_issues, "shared_formula_groups": topology.get("shared_formula_groups", []),
            "xml_formula_gaps": topology.get("gaps", []), "cells": cells}


def extract(path: str | Path) -> dict[str, Any]:
    path = Path(path)
    workbook = openpyxl.load_workbook(path, data_only=False)
    topology = _formula_topology(path)
    flow_sheet = next((sheet for sheet in workbook.worksheets if "流程" in sheet.title), None)
    template_sheets = [sheet for sheet in workbook.worksheets if sheet is not flow_sheet and "考核指标" in {cell.value for row in sheet.iter_rows() for cell in row}]
    hr_sheets = [sheet for sheet in workbook.worksheets if "HR汇总" in sheet.title]
    cases = [_extract_source_data(sheet, index) for index, sheet in enumerate(template_sheets)]
    hr = [_parse_hr(sheet, topology.get(sheet.title, {})) for sheet in hr_sheets]
    hr_names = [entry["sheet"] for entry in hr]
    return {"source": str(path), "source_filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "workbook_sheet_count": len(workbook.worksheets), "template_sheet_count": len(template_sheets),
            "template_sheets": [sheet.title for sheet in template_sheets], "source_department_count": len({case["department"] for case in cases}),
            "source_departments": sorted({case["department"] for case in cases}),
            "score_policy": "Preserve source numeric H scores; blank H remains unscored. Blank, zero, and not-applicable remain distinct in source metadata.",
            "rounding": "ROUND dimension*item/100 to 1 decimal, ROUND contribution to 2 decimals, SUM contributions.",
            "flow": _parse_flow(flow_sheet), "cases": cases, "hr": hr,
            "hr_source_cells": {entry["sheet"]: entry["cells"] for entry in hr},
            "formula_diagnostics": {name: {"shared_formula_groups": data.get("shared_formula_groups", []), "xml_formula_gaps": data.get("gaps", [])}
                                    for name, data in topology.items() if name in hr_names}}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Verify and extract a Hongxing source workbook without modifying it")
    parser.add_argument("workbook")
    parser.add_argument("--output", help="Write JSON diagnostics to this path instead of stdout")
    args = parser.parse_args()
    result = json.dumps(extract(args.workbook), ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(result + "\n", encoding="utf-8")
    else:
        print(result)
